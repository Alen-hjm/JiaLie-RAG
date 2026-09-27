import re
from pathlib import Path
from docx import Document
from pypdf import PdfReader


class DocumentParseError(ValueError):
    pass


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            reader = PdfReader(str(path))
            pages = [(page.extract_text() or "").strip() for page in reader.pages]
            text = "\n\n".join(f"[第 {i + 1} 页]\n{content}" for i, content in enumerate(pages) if content)
        elif suffix == ".docx":
            doc = Document(str(path))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                paragraphs.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
            text = "\n".join(paragraphs)
        else:
            raise DocumentParseError("仅支持 PDF 和 DOCX 文件")
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError(f"文档解析失败：{exc}") from exc
    if len(re.sub(r"\s", "", text)) < 30:
        raise DocumentParseError("未提取到足够文字；扫描版 PDF 暂不支持，请先进行 OCR")
    return text


HEADING_PATTERN = r"工作经历|项目经历|教育经历|专业技能|个人总结"
HEADING_RE = re.compile(rf"^(?:{HEADING_PATTERN})")


def split_chunks(text: str, max_chars: int = 900, overlap: int = 120) -> list[dict]:
    """Split extracted resume text into retrievable chunks.

    Section headings act as *hard* boundaries: previously every short resume
    collapsed into a single chunk (sections were merged until the 900-character
    budget was hit), which meant the whole resume shared one embedding, the
    evidence list could only ever show one entry, and "切分" was effectively a
    no-op on a normal one-page resume.

    Oversized sections are still hard-split with an overlap so no content is
    lost across the boundary.
    """
    sections = re.split(rf"\n{{2,}}|(?={HEADING_PATTERN})", text)
    chunks: list[dict] = []
    buffer = ""

    def flush() -> None:
        nonlocal buffer
        buffer = buffer.strip()
        if buffer:
            chunks.append({"content": buffer, "section": infer_section(buffer), "page_number": infer_page(buffer)})
        buffer = ""

    for section in sections:
        section = section.strip()
        if not section:
            continue
        if HEADING_RE.match(section) and buffer:
            flush()
        if len(buffer) + len(section) + 1 <= max_chars:
            buffer = f"{buffer}\n{section}".strip()
            continue
        flush()
        buffer = section
        while len(buffer) > max_chars:
            piece = buffer[:max_chars]
            chunks.append({"content": piece, "section": infer_section(piece), "page_number": infer_page(piece)})
            buffer = buffer[max_chars - overlap:]
    flush()
    return chunks


def infer_section(text: str) -> str:
    for marker, name in [("工作", "工作经历"), ("项目", "项目经历"), ("教育", "教育经历"), ("技能", "专业技能")]:
        if marker in text[:80]:
            return name
    return "简历正文"


def infer_page(text: str) -> int | None:
    match = re.search(r"\[第\s*(\d+)\s*页\]", text)
    return int(match.group(1)) if match else None

