from pathlib import Path
from docx import Document
from app.services.documents import DocumentParseError, extract_text, split_chunks


def test_extract_docx_and_chunks(tmp_path: Path):
    path = tmp_path / "resume.docx"
    doc = Document(); doc.add_paragraph("林若川"); doc.add_paragraph("工作经历"); doc.add_paragraph("十年半导体销售经验，负责晶圆厂大客户与团队管理。")
    doc.save(path)
    text = extract_text(path)
    assert "半导体销售" in text
    assert split_chunks(text)[0]["section"] in {"简历正文", "工作经历"}


def test_rejects_unknown_format(tmp_path: Path):
    path = tmp_path / "resume.txt"; path.write_text("简历" * 30, encoding="utf-8")
    try:
        extract_text(path)
        assert False
    except DocumentParseError as exc:
        assert "仅支持" in str(exc)

